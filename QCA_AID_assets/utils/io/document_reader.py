"""
Document Reader

Reads and processes various document formats for QCA-AID.
Supports TXT, DOCX, and PDF files with text extraction and cleaning.
"""

import os
from pathlib import Path
from typing import Optional, Dict, List, Tuple
import re

from ..export.helpers import sanitize_text_for_excel


try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DocxDocument = None
    DOCX_AVAILABLE = False

try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PyPDF2 = None
    PDF_AVAILABLE = False


class DocumentReader:
    def __init__(self, data_dir: str):
        self.data_dir = data_dir
        self.supported_formats = {'.docx', '.pdf', '.txt'}
        os.makedirs(data_dir, exist_ok=True)
        
        # Prüfe verfügbare Features
        if not DOCX_AVAILABLE:
            print("⚠️ python-docx nicht verfügbar - DOCX-Unterstützung deaktiviert")
            self.supported_formats.discard('.docx')
        
        if not PDF_AVAILABLE:
            print("⚠️ PyPDF2 nicht verfügbar - PDF-Unterstützung deaktiviert")
            self.supported_formats.discard('.pdf')
        
        print(f"\nDocumentReader initialisiert:")
        print(f"Verzeichnis: {os.path.abspath(data_dir)}")
        print(f"Unterstützte Formate: {', '.join(self.supported_formats)}")


    def clean_problematic_characters(self, text: str) -> str:
        """Verwendet die bereits vorhandene Funktion"""
        return sanitize_text_for_excel(text)

    async def read_documents(self) -> Dict[str, str]:
        documents = {}
        try:
            all_files = os.listdir(self.data_dir)
            
            print("\nDateianalyse:")
            def is_supported_file(filename: str) -> bool:
                # Exclude backup and temporary files
                if any(ext in filename.lower() for ext in ['.bak', '.bkk', '.tmp', '~']):
                    return False
                # Get the file extension
                extension = os.path.splitext(filename)[1].lower()
                return extension in self.supported_formats

            supported_files = [f for f in all_files if is_supported_file(f)]
            
            print(f"\nGefundene Dateien:")
            for file in all_files:
                status = "[OK]" if is_supported_file(file) else "[ERROR]"
                print(f"{status} {file}")
            
            print(f"\nVerarbeite Dateien:")
            for filename in supported_files:
                try:
                    filepath = os.path.join(self.data_dir, filename)
                    extension = os.path.splitext(filename)[1].lower()
                    
                    print(f"\nLese: {filename}")
                    
                    if extension == '.docx':
                        content = self._read_docx(filepath)
                    elif extension == '.pdf':
                        content = self._read_pdf(filepath)
                    elif extension == '.txt':
                        content = self._read_txt(filepath)
                    else:
                        print(f"⚠ Nicht unterstütztes Format: {extension}")
                        continue
                    
                    if content and content.strip():
                        documents[filename] = content
                        print(f"[OK] Erfolgreich eingelesen: {len(content)} Zeichen")
                    else:
                        print(f"⚠ Keine Textinhalte gefunden")
                
                except Exception as e:
                    print(f"[ERROR] Fehler bei {filename}: {str(e)}")
                    print("Details:")
                    import traceback
                    traceback.print_exc()
                    continue

            print(f"\nVerarbeitungsstatistik:")
            print(f"- Dateien im Verzeichnis: {len(all_files)}")
            print(f"- Unterstützte Dateien: {len(supported_files)}")
            print(f"- Erfolgreich eingelesen: {len(documents)}")
            
            return documents

        except Exception as e:
            print(f"Fehler beim Einlesen der Dokumente: {str(e)}")
            import traceback
            traceback.print_exc()
            return documents

    def _read_txt(self, filepath: str) -> str:
        """Liest eine Textdatei ein."""
        with open(filepath, 'r', encoding='utf-8') as f:
            return f.read()

    def _read_docx(self, filepath: str) -> str:
        """
        Liest eine DOCX-Datei ein und extrahiert den Text mit ausführlicher Diagnose.
        Enthält zusätzliche Bereinigung für problematische Zeichen.
        """
        try:
            from docx import Document
            print(f"\nDetailierte Analyse von: {os.path.basename(filepath)}")
            
            # Öffne das Dokument mit zusätzlicher Fehlerbehandlung
            try:
                doc = Document(filepath)
            except Exception as e:
                print(f"  Fehler beim Öffnen der Datei: {str(e)}")
                print("  Versuche alternative Öffnungsmethode...")
                # Manchmal hilft es, die Datei zuerst zu kopieren
                import shutil
                temp_path = filepath + '.temp'
                shutil.copy2(filepath, temp_path)
                doc = Document(temp_path)
                os.remove(temp_path)

            # Sammle Dokumentinformationen
            paragraphs = []
            print("\nDokumentanalyse:")
            print(f"  Gefundene Paragraphen: {len(doc.paragraphs)}")
            
            # Verarbeite jeden Paragraphen einzeln
            for i, paragraph in enumerate(doc.paragraphs):
                text = paragraph.text.strip()
                if text:
                    # Hier bereinigen wir Steuerzeichen und problematische Zeichen
                    clean_text = self.clean_problematic_characters(text)
                    paragraphs.append(clean_text)
                else:
                    print(f"  Paragraph {i+1}: Leer")

            # Wenn Paragraphen gefunden wurden
            if paragraphs:
                full_text = '\n'.join(paragraphs)
                print(f"\nErgebnis:")
                print(f"  [OK] {len(paragraphs)} Textparagraphen extrahiert")
                print(f"  [OK] Gesamtlänge: {len(full_text)} Zeichen")
                return full_text
            
            # Wenn keine Paragraphen gefunden wurden, suche in anderen Bereichen
            print("\nSuche nach alternativen Textinhalten:")
            
            # Prüfe Tabellen
            table_texts = []
            for i, table in enumerate(doc.tables):
                print(f"  Prüfe Tabelle {i+1}:")
                for row in table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            # Auch hier bereinigen
                            clean_text = self.clean_problematic_characters(cell_text)
                            table_texts.append(clean_text)
                            print(f"    Zelleninhalt gefunden: {len(cell_text)} Zeichen")
            
            if table_texts:
                full_text = '\n'.join(table_texts)
                print(f"\nErgebnis:")
                print(f"  [OK] {len(table_texts)} Tabelleneinträge extrahiert")
                print(f"  [OK] Gesamtlänge: {len(full_text)} Zeichen")
                return full_text
                
            print("\n[ERROR] Keine Textinhalte im Dokument gefunden")
            return ""
                
        except ImportError:
            print("\n[ERROR] python-docx nicht installiert.")
            print("  Bitte installieren Sie das Paket mit:")
            print("  pip install python-docx")
            raise
        except Exception as e:
            print(f"\n[ERROR] Unerwarteter Fehler beim DOCX-Lesen:")
            print(f"  {str(e)}")
            import traceback
            traceback.print_exc()
            raise

    def _read_pdf(self, filepath: str) -> str:
        """
        Liest eine PDF-Datei ein und extrahiert den Text mit verbesserter Fehlerbehandlung.
        
        Args:
            filepath: Pfad zur PDF-Datei
                
        Returns:
            str: Extrahierter und bereinigter Text
        """
        try:
            import PyPDF2
            print(f"\nLese PDF: {os.path.basename(filepath)}")
            
            text_content = []
            with open(filepath, 'rb') as file:
                try:
                    pdf_reader = PyPDF2.PdfReader(file)
                    total_pages = len(pdf_reader.pages)
                    print(f"  Gefundene Seiten: {total_pages}")
                    
                    for page_num, page in enumerate(pdf_reader.pages, 1):
                        try:
                            page_text = page.extract_text()
                            if page_text:
                                # Bereinige den Text von problematischen Zeichen
                                cleaned_text = self.clean_problematic_characters(page_text)
                                text_content.append(cleaned_text)
                                print(f"  Seite {page_num}/{total_pages}: {len(cleaned_text)} Zeichen extrahiert")
                            else:
                                print(f"  Seite {page_num}/{total_pages}: Kein Text gefunden")
                                
                        except Exception as page_error:
                            print(f"  Fehler bei Seite {page_num}: {str(page_error)}")
                            # Versuche es mit alternativer Methode
                            try:
                                # Fallback: Extrahiere einzelne Textfragmente
                                print("  Versuche alternative Extraktionsmethode...")
                                if hasattr(page, 'extract_text'):
                                    fragments = []
                                    for obj in page.get_text_extraction_elements():
                                        if hasattr(obj, 'get_text'):
                                            fragments.append(obj.get_text())
                                    if fragments:
                                        fallback_text = ' '.join(fragments)
                                        cleaned_text = self.clean_problematic_characters(fallback_text)
                                        text_content.append(cleaned_text)
                                        print(f"  Alternative Methode erfolgreich: {len(cleaned_text)} Zeichen")
                            except:
                                print("  Alternative Methode fehlgeschlagen")
                                continue
                    
                except PyPDF2.errors.PdfReadError as pdf_error:
                    print(f"  PDF Lesefehler: {str(pdf_error)}")
                    print("  Versuche Fallback-Methode...")
                    
                    # Fallback-Methode, wenn direkte Lesemethode fehlschlägt
                    try:
                        from pdf2image import convert_from_path
                        from pytesseract import image_to_string
                        
                        print("  Verwende OCR-Fallback via pdf2image und pytesseract")
                        # Konvertiere PDF-Seiten zu Bildern
                        images = convert_from_path(filepath)
                        
                        for i, image in enumerate(images):
                            try:
                                # Extrahiere Text über OCR
                                ocr_text = image_to_string(image, lang='deu')
                                if ocr_text:
                                    # Bereinige den Text
                                    cleaned_text = self.clean_problematic_characters(ocr_text)
                                    text_content.append(cleaned_text)
                                    print(f"  OCR Seite {i+1}: {len(cleaned_text)} Zeichen extrahiert")
                                else:
                                    print(f"  OCR Seite {i+1}: Kein Text gefunden")
                            except Exception as ocr_error:
                                print(f"  OCR-Fehler bei Seite {i+1}: {str(ocr_error)}")
                                continue
                    
                    except ImportError:
                        print("  OCR-Fallback nicht verfügbar. Bitte installieren Sie pdf2image und pytesseract")
                    
            # Zusammenfassen des extrahierten Textes
            if text_content:
                full_text = '\n'.join(text_content)
                print(f"\nErgebnis:")
                print(f"  [OK] {len(text_content)} Textabschnitte extrahiert")
                print(f"  [OK] Gesamtlänge: {len(full_text)} Zeichen")
                return full_text
            else:
                print("\n[ERROR] Kein Text aus PDF extrahiert")
                return ""
                
        except ImportError:
            print("PyPDF2 nicht installiert. Bitte installieren Sie: pip install PyPDF2")
            raise
        except Exception as e:
            print(f"Fehler beim PDF-Lesen: {str(e)}")
            import traceback
            traceback.print_exc()
            return ""  # Leerer String im Fehlerfall, damit der Rest funktioniert
    
    def _extract_metadata(self, filename: str) -> Tuple[str, str, str]:
        """
        Extrahiert Metadaten aus dem Dateinamen.
        Erwartet Format: attribut1_attribut2_attribut3.extension
        
        Args:
            filename (str): Name der Datei
            
        Returns:
            Tuple[str, str, str]: (attribut1, attribut2, attribut3)
        """
        try:
            name_without_ext = os.path.splitext(filename)[0]
            parts = name_without_ext.split('_')
            
            # Extrahiere bis zu drei Attribute, wenn verfügbar
            attribut1 = parts[0] if len(parts) >= 1 else ""
            attribut2 = parts[1] if len(parts) >= 2 else ""
            attribut3 = parts[2] if len(parts) >= 3 else ""
            
            return attribut1, attribut2, attribut3
        except Exception as e:
            print(f"Fehler beim Extrahieren der Metadaten aus {filename}: {str(e)}")
            return filename, "", ""

